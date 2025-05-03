import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime

projects_data = pd.DataFrame({
    'Project ID': [101, 102, 103],
    'Project Name': ['Website Redesign', 'Marketing Campaign', 'Product Launch'],
    'Start Date': [datetime(2025, 1, 10), datetime(2025, 2, 1), datetime(2025, 3, 5)],
    'End Date': [datetime(2025, 4, 10), datetime(2025, 4, 30), datetime(2025, 6, 1)],
    'Status': ['In Progress', 'Planned', 'Not Started'],
    'Manager': ['Alice', 'Bob', 'Charlie']
})

budget_data = pd.DataFrame({
    'Project ID': [101, 101, 102, 103],
    'Category': ['Design', 'Development', 'Ad Spend', 'Production'],
    'Forecasted Budget': [5000, 10000, 7000, 12000],
    'Actual Spend': [4500, 9000, 0, 0]
})

wb = Workbook()
ws_projects = wb.active
ws_projects.title = "Projects"

ws_budget = wb.create_sheet("Budget")
ws_summary = wb.create_sheet("Summary")

for r in dataframe_to_rows(projects_data, index=False, header=True):
    ws_projects.append(r)

for cell in ws_projects[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="D9EAD3")
    cell.alignment = Alignment(horizontal="center")

for r in dataframe_to_rows(budget_data, index=False, header=True):
    ws_budget.append(r)

for cell in ws_budget[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="F9CB9C")
    cell.alignment = Alignment(horizontal="center")

summary_data = budget_data.copy()
summary_data = summary_data.groupby("Project ID").agg({
    "Forecasted Budget": "sum",
    "Actual Spend": "sum"
}).reset_index()

summary_data = summary_data.merge(projects_data[["Project ID", "Project Name"]], on="Project ID")
summary_data["Remaining Budget"] = summary_data["Forecasted Budget"] - summary_data["Actual Spend"]
summary_data = summary_data[[
    "Project ID", "Project Name", "Forecasted Budget", "Actual Spend", "Remaining Budget"
]]

for r in dataframe_to_rows(summary_data, index=False, header=True):
    ws_summary.append(r)

for cell in ws_summary[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="CFE2F3")
    cell.alignment = Alignment(horizontal="center")

for ws in [ws_projects, ws_budget, ws_summary]:
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

wb.save("Project_Budget_Tracker_2025.xlsx")
print("✅ Excel file 'Project_Budget_Tracker_2025.xlsx' created successfully!")
